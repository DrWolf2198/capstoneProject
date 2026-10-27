"""Laporan perkembangan anak: gabungan penilaian, absensi, dokumentasi, catatan."""
from datetime import date, timedelta
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func
from sqlalchemy.orm import Session

from .. import access, models
from ..database import get_db
from ..deps import get_current_user

router = APIRouter()


def _parse_bulan(bulan: str | None) -> tuple[date, date] | None:
    """Parse YYYY-MM -> (awal_bulan, awal_bulan_berikutnya). None jika tidak ada."""
    if not bulan:
        return None
    try:
        y, m = bulan.split("-")
        y, m = int(y), int(m)
        start = date(y, m, 1)
        end = date(y, m + 1, 1) if m < 12 else date(y + 1, 1, 1)
        return start, end
    except (ValueError, IndexError):
        raise HTTPException(400, "Format bulan tidak valid (YYYY-MM)")


@router.get("/murid/{murid_id}")
def laporan_murid(
    murid_id: int,
    bulan: str | None = None,
    db: Session = Depends(get_db),
    user: models.Pengguna = Depends(get_current_user),
):
    """Laporan perkembangan satu murid (opsional filter bulan YYYY-MM)."""
    if not access.can_access_murid(user, murid_id, db):
        raise HTTPException(403, "Akses ditolak")

    m = db.get(models.Murid, murid_id)
    if m is None:
        raise HTTPException(404, "Murid tidak ditemukan")

    rentang = _parse_bulan(bulan)

    # --- penilaian: pencatatan terbaru per aspek (dalam periode) ---
    q_pen = db.query(models.PencatatanPerkembangan).filter(
        models.PencatatanPerkembangan.murid_id == murid_id
    )
    if rentang:
        q_pen = q_pen.filter(
            models.PencatatanPerkembangan.tanggal >= rentang[0],
            models.PencatatanPerkembangan.tanggal < rentang[1],
        )
    per_aspek: dict[int, models.PencatatanPerkembangan] = {}
    for p in q_pen.order_by(models.PencatatanPerkembangan.tanggal.asc()).all():
        per_aspek[p.aspek_id] = p  # terakhir menang (asc, overwrite -> latest)
    penilaian = [
        {
            "aspek": p.aspek.nama if p.aspek else None,
            "kategori": p.kategori,
            "deskripsi": p.deskripsi,
            "tanggal": p.tanggal.isoformat(),
        }
        for p in per_aspek.values()
    ]

    # --- absensi rekap (periode) ---
    q_abs = (
        db.query(models.Absensi)
        .join(models.SesiKelas, models.Absensi.sesi_id == models.SesiKelas.id)
        .filter(models.Absensi.murid_id == murid_id)
    )
    if rentang:
        q_abs = q_abs.filter(
            models.SesiKelas.tanggal >= rentang[0],
            models.SesiKelas.tanggal < rentang[1],
        )
    rows = q_abs.all()
    hadir = sum(1 for r in rows if r.status == "hadir")
    izin = sum(1 for r in rows if r.status == "izin")
    sakit = sum(1 for r in rows if r.status == "sakit")
    alfa = sum(1 for r in rows if r.status == "alfa")
    total = len(rows)
    absensi = {
        "hadir": hadir,
        "izin": izin,
        "sakit": sakit,
        "alfa": alfa,
        "total": total,
        "persentase": round(hadir / total * 100, 1) if total else None,
    }

    # --- dokumentasi (periode) ---
    q_dok = db.query(models.Dokumentasi).filter(models.Dokumentasi.murid_id == murid_id)
    if rentang:
        q_dok = q_dok.filter(
            models.Dokumentasi.tanggal >= rentang[0],
            models.Dokumentasi.tanggal < rentang[1],
        )
    dokumentasi = [
        {
            "tipe": d.tipe,
            "url": d.url,
            "deskripsi": d.deskripsi,
            "tanggal": d.tanggal.isoformat(),
        }
        for d in q_dok.order_by(models.Dokumentasi.tanggal.desc()).all()
    ]

    # --- catatan (periode) ---
    q_cat = db.query(models.Catatan).filter(models.Catatan.murid_id == murid_id)
    if rentang:
        q_cat = q_cat.filter(
            models.Catatan.created_at >= rentang[0],
            models.Catatan.created_at < rentang[1],
        )
    catatan = [
        {
            "isi": c.isi,
            "nama_pengguna": c.pengguna.nama if c.pengguna else None,
            "role": c.pengguna.role if c.pengguna else None,
            "created_at": c.created_at.isoformat() if c.created_at else None,
        }
        for c in q_cat.order_by(models.Catatan.created_at.asc()).all()
    ]

    return {
        "murid": {
            "id": m.id,
            "nama": m.nama,
            "tanggal_lahir": m.tanggal_lahir.isoformat() if m.tanggal_lahir else None,
            "jenis_kelamin": m.jenis_kelamin,
        },
        "kelas": m.kelas.nama if m.kelas else None,
        "bulan": bulan,
        "penilaian": penilaian,
        "absensi": absensi,
        "dokumentasi": dokumentasi,
        "catatan": catatan,
    }


@router.get("/export/penilaian")
def export_penilaian(
    kelas_id: int | None = None,
    bulan: str | None = None,
    db: Session = Depends(get_db),
    user: models.Pengguna = Depends(get_current_user),
):
    """Export data penilaian ke Excel (.xlsx) — scoped ke murid yang bisa diakses user."""
    ids = access.murid_ids_for(user, db)
    if not ids:
        raise HTTPException(403, "Tidak ada akses")

    rentang = _parse_bulan(bulan)

    q = (
        db.query(models.PencatatanPerkembangan)
        .join(models.Murid, models.PencatatanPerkembangan.murid_id == models.Murid.id)
        .filter(models.Murid.id.in_(ids))
    )
    if kelas_id:
        q = q.filter(models.Murid.kelas_id == kelas_id)
    if rentang:
        q = q.filter(
            models.PencatatanPerkembangan.tanggal >= rentang[0],
            models.PencatatanPerkembangan.tanggal < rentang[1],
        )

    rows = q.order_by(models.Murid.nama, models.PencatatanPerkembangan.tanggal).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Penilaian"
    ws.append(["No", "Nama Murid", "Kelas", "Aspek", "Capaian", "Catatan", "Tanggal"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, p in enumerate(rows, 1):
        ws.append([
            i,
            p.murid.nama,
            p.murid.kelas.nama if p.murid.kelas else "-",
            p.aspek.nama if p.aspek else "-",
            p.kategori,
            p.deskripsi,
            p.tanggal.isoformat(),
        ])

    for idx, w in enumerate([5, 26, 10, 24, 10, 42, 12], 1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"penilaian-{bulan if bulan else 'semua'}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/export/absensi")
def export_absensi(
    murid_ids: str,
    periode: int = 1,
    db: Session = Depends(get_db),
    user: models.Pengguna = Depends(get_current_user),
):
    """Export rekap kehadiran (multi murid + periode) ke Excel (.xlsx)."""
    ids = [int(x) for x in murid_ids.split(",") if x.strip().isdigit()]
    if not ids:
        raise HTTPException(400, "Pilih minimal satu murid")
    accessible = set(access.murid_ids_for(user, db))
    ids = [i for i in ids if i in accessible]
    if not ids:
        raise HTTPException(403, "Tidak ada akses ke murid terpilih")
    if periode not in (1, 3, 6, 12):
        raise HTTPException(400, "Periode tidak valid (1/3/6/12 bulan)")

    start = date.today() - timedelta(days=30 * periode)

    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap Kehadiran"
    ws.append(["No", "Nama Murid", "Kelas", "Hadir", "Izin", "Sakit", "Alfa", "Total Pertemuan", "Kehadiran (%)"])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    murid_rows = (
        db.query(models.Murid)
        .filter(models.Murid.id.in_(ids))
        .order_by(models.Murid.nama)
        .all()
    )

    # tanggal libur dalam periode (dari kalender akademik) — tidak dihitung sebagai pertemuan
    libur_dates = {
        d[0]
        for d in db.query(models.AgendaAkademik.tanggal)
        .filter(
            models.AgendaAkademik.tanggal >= start,
            models.AgendaAkademik.kategori.ilike("%libur%"),
        )
        .all()
    }

    # total pertemuan (hari unik) per kelas dalam periode — sesi di tanggal libur dikecualikan
    kelas_ids = {m.kelas_id for m in murid_rows if m.kelas_id}
    total_sesi_map: dict[int, int] = {}
    for kid in kelas_ids:
        q = db.query(func.count(func.distinct(models.SesiKelas.tanggal))).filter(
            models.SesiKelas.kelas_id == kid,
            models.SesiKelas.tanggal >= start,
        )
        if libur_dates:
            q = q.filter(models.SesiKelas.tanggal.notin_(libur_dates))
        total_sesi_map[kid] = q.scalar() or 0

    for i, m in enumerate(murid_rows, 1):
        q_rows = (
            db.query(models.Absensi)
            .join(models.SesiKelas, models.Absensi.sesi_id == models.SesiKelas.id)
            .filter(models.Absensi.murid_id == m.id, models.SesiKelas.tanggal >= start)
        )
        if libur_dates:
            q_rows = q_rows.filter(models.SesiKelas.tanggal.notin_(libur_dates))
        rows = q_rows.all()
        hadir = sum(1 for r in rows if r.status == "hadir")
        izin = sum(1 for r in rows if r.status == "izin")
        sakit = sum(1 for r in rows if r.status == "sakit")
        alfa = sum(1 for r in rows if r.status == "alfa")
        total_sesi = total_sesi_map.get(m.kelas_id, 0)
        persen = round(hadir / total_sesi * 100, 1) if total_sesi else 0
        ws.append([
            i,
            m.nama,
            m.kelas.nama if m.kelas else "-",
            hadir,
            izin,
            sakit,
            alfa,
            total_sesi,
            persen,
        ])

    for idx, w in enumerate([5, 26, 12, 8, 8, 8, 8, 14, 14], 1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"rekap-kehadiran-{periode}bulan.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
